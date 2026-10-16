"""Bounded, isolated document rendering. Never returns original document bytes."""
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile

MAX_FILE = 10 * 1024 * 1024
MAX_PAGES = 30


def render_preview(content: bytes, extension: str, page: int):
    if not content or len(content) > MAX_FILE or extension not in ('.pdf', '.docx') or page < 0:
        raise ValueError('Unsupported or oversized resume')
    with tempfile.TemporaryDirectory(prefix='resume-preview-') as folder:
        root = Path(folder)
        source = root / ('resume' + extension)
        source.write_bytes(content)
        subprocess.run([sys.executable, str(Path(__file__).resolve()), str(source), str(page), str(root)],
                       check=True, timeout=15, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        meta = json.loads((root / 'meta.json').read_text())
        return meta, (root / 'page.png').read_bytes() if page else None


def worker(source, page, root):
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (512 * 1024**2, 512 * 1024**2))
    resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
    resource.setrlimit(resource.RLIMIT_FSIZE, (20 * 1024**2, 20 * 1024**2))
    import fitz
    if source.suffix == '.pdf':
        document = fitz.open(source)
    else:
        # DOCX preview preserves text/tables; it deliberately excludes embedded
        # objects, scripts, relationships and remotely referenced resources.
        import zipfile
        with zipfile.ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024**2:
                raise ValueError('Expanded document too large')
        from docx import Document
        from html import escape
        doc = Document(source)
        html = ''.join('<p>' + escape(p.text) + '</p>' for p in doc.paragraphs)
        html += ''.join('<table>' + ''.join('<tr>' + ''.join('<td>' + escape(c.text) + '</td>' for c in row.cells) + '</tr>' for row in table.rows) + '</table>' for table in doc.tables)
        output = io.BytesIO()
        writer = fitz.DocumentWriter(output)
        story = fitz.Story(html=html, user_css='body {font-family: sans-serif; font-size: 11pt;} td {padding: 4pt;}')
        more = True
        count = 0
        while more:
            count += 1
            if count > MAX_PAGES:
                raise ValueError('Too many pages')
            device = writer.begin_page(fitz.paper_rect('a4'))
            more, _ = story.place(fitz.Rect(36, 36, 559, 806))
            story.draw(device)
            writer.end_page()
        writer.close()
        document = fitz.open(stream=output.getvalue(), filetype='pdf')
    if document.needs_pass or not 0 < document.page_count <= MAX_PAGES or page > document.page_count:
        raise ValueError('Document cannot be displayed')
    (root / 'meta.json').write_text(json.dumps({'pages': document.page_count}))
    if page:
        current = document[page - 1]
        if current.rect.width <= 0 or current.rect.height <= 0:
            raise ValueError('Invalid page')
        scale = min(1.5, 1600 / max(current.rect.width, current.rect.height))
        pix = current.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pix.save(root / 'page.png')
    document.close()


if __name__ == '__main__':
    worker(Path(sys.argv[1]), int(sys.argv[2]), Path(sys.argv[3]))
