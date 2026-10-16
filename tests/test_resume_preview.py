import io
import unittest
import fitz
from docx import Document
from app.resume_preview import render_preview

class ResumePreviewTests(unittest.TestCase):
    def pdf(self, pages=2):
        doc = fitz.open()
        for i in range(pages):
            doc.new_page().insert_text((60, 60), f'Resume page {i+1}')
        data = doc.tobytes()
        doc.close()
        return data

    def test_pdf_metadata_and_page_are_separate(self):
        meta, image = render_preview(self.pdf(), '.pdf', 0)
        self.assertEqual(meta, {'pages': 2})
        self.assertIsNone(image)
        _, image = render_preview(self.pdf(), '.pdf', 2)
        self.assertTrue(image.startswith(b'\x89PNG\r\n\x1a\n'))
        self.assertNotIn(b'%PDF', image[:16])

    def test_docx_text_and_table(self):
        doc = Document()
        doc.add_paragraph('Resume <script>escaped content</script>')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Python / AWS'
        stream = io.BytesIO()
        doc.save(stream)
        meta, image = render_preview(stream.getvalue(), '.docx', 1)
        self.assertEqual(meta['pages'], 1)
        self.assertTrue(image.startswith(b'\x89PNG'))

    def test_rejects_unsupported_oversized_and_invalid_files(self):
        for content, ext, page in [(b'x', '.doc', 0), (b'x', '.pdf', -1), (b'x' * (10*1024**2+1), '.pdf', 0)]:
            with self.assertRaises(ValueError): render_preview(content, ext, page)
        for content, page in [(b'invalid', 0), (self.pdf(31), 0), (self.pdf(), 3)]:
            with self.assertRaises(Exception): render_preview(content, '.pdf', page)

if __name__ == '__main__': unittest.main()
